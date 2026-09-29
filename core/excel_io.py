from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from core.company import (
    COMPANY_ADDRESS,
    COMPANY_NAME,
    COMPANY_TEL,
    DOCUMENT_TITLE,
    HEADER_ROWS,
    LOGO_PATH,
)
from core.invoice import Invoice, ProductRow, parse_number

THIN = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)
HEADER_FILL = PatternFill("solid", fgColor="217346")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Times New Roman", bold=True, size=20)
COMPANY_FONT = Font(name="Times New Roman", bold=True, size=14)
ADDR_FONT = Font(name="Times New Roman", size=11)
LABEL_FONT = Font(name="Times New Roman", bold=True, size=12)
CELL_FONT = Font(name="Calibri", size=11)
RED_FONT = Font(name="Calibri", size=11, color="C00000")
LOCKED_FILL = PatternFill("solid", fgColor="F2F2F2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def _style_header_cell(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN


def export_invoice(invoice: Invoice, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoa don"

    last_col = invoice.col_count()
    last_letter = get_column_letter(last_col)

    # Logo & Company Header
    ws.merge_cells("A1:B4")
    ws.merge_cells(f"C1:{last_letter}1")
    ws.merge_cells(f"C2:{last_letter}2")
    ws.merge_cells(f"C3:{last_letter}3")
    ws.merge_cells(f"A6:{last_letter}7")

    ws["C1"] = COMPANY_NAME
    ws["C1"].font = Font(name="Times New Roman", bold=True, size=13)
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["C2"] = COMPANY_ADDRESS
    ws["C2"].font = ADDR_FONT
    ws["C2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["C3"] = COMPANY_TEL
    ws["C3"].font = ADDR_FONT
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A6"] = DOCUMENT_TITLE
    ws["A6"].font = TITLE_FONT
    ws["A6"].alignment = Alignment(horizontal="center", vertical="center")

    if LOGO_PATH.exists():
        img = XLImage(str(LOGO_PATH))
        img.width = 118
        img.height = 74
        img.anchor = "A1"
        ws.add_image(img)

    # Customer Info Block Layout
    if last_col <= 7:
        date_col = 6
        name_end_col = 5
    else:
        date_col = max(6, last_col - 2)
        name_end_col = date_col - 1

    date_letter = get_column_letter(date_col)
    name_end_letter = get_column_letter(name_end_col)
    date_val_start_letter = get_column_letter(date_col + 1)

    # Row 9-10: Tên khách & Ngày
    ws.merge_cells("A9:B10")
    ws["A9"] = "Tên khách:"
    ws["A9"].font = LABEL_FONT
    ws["A9"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(f"C9:{name_end_letter}10")
    name_cell = ws["C9"]
    name_cell.value = invoice.customer_name
    name_cell.font = ADDR_FONT
    name_cell.alignment = Alignment(horizontal="left", vertical="center")

    for c in range(3, name_end_col + 1):
        ws.cell(10, c).border = Border(bottom=Side(style="thin"))

    ws.merge_cells(f"{date_letter}9:{date_letter}10")
    date_label_cell = ws[f"{date_letter}9"]
    date_label_cell.value = "Ngày:"
    date_label_cell.font = LABEL_FONT
    date_label_cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(f"{date_val_start_letter}9:{last_letter}10")
    date_cell = ws[f"{date_val_start_letter}9"]
    date_cell.value = invoice.invoice_date
    date_cell.number_format = "DD/MM/YYYY"
    date_cell.font = ADDR_FONT
    date_cell.alignment = CENTER

    for c in range(date_col + 1, last_col + 1):
        ws.cell(10, c).border = Border(bottom=Side(style="thin"))

    # Row 12-13: Địa chỉ & Số
    ws.merge_cells("A12:B13")
    ws["A12"] = "Địa chỉ:"
    ws["A12"].font = LABEL_FONT
    ws["A12"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(f"C12:{name_end_letter}13")
    addr_cell = ws["C12"]
    addr_cell.value = invoice.customer_address
    addr_cell.font = ADDR_FONT
    addr_cell.alignment = Alignment(horizontal="left", vertical="center")

    for c in range(3, name_end_col + 1):
        ws.cell(13, c).border = Border(bottom=Side(style="thin"))

    ws.merge_cells(f"{date_letter}12:{date_letter}13")
    num_label_cell = ws[f"{date_letter}12"]
    num_label_cell.value = "Số:"
    num_label_cell.font = LABEL_FONT
    num_label_cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(f"{date_val_start_letter}12:{last_letter}13")
    num_cell = ws[f"{date_val_start_letter}12"]
    num_cell.value = invoice.invoice_number
    num_cell.font = ADDR_FONT
    num_cell.alignment = CENTER

    for c in range(date_col + 1, last_col + 1):
        ws.cell(13, c).border = Border(bottom=Side(style="thin"))

    header_row = HEADER_ROWS
    labels = invoice.header_labels()
    for i, label in enumerate(labels, start=1):
        cell = ws.cell(header_row, i, label)
        _style_header_cell(cell)

    widths = [6, 16, 32, 10, 12]
    widths.extend([9] * invoice.n_xe)
    widths.append(12)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 16

    for r_idx, row in enumerate(invoice.rows):
        excel_row = HEADER_ROWS + 1 + r_idx
        values = [
            r_idx + 1,
            row.ma,
            row.ten,
            row.dvt,
            row.so_luong if row.so_luong else 0,
        ]
        for v in row.xe:
            values.append(v if v not in (None, "") else None)
        values.append(None)

        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(excel_row, c_idx, val)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = CENTER
            if c_idx == 3:
                cell.alignment = LEFT
            if c_idx in (1, 5) or c_idx >= 6:
                if isinstance(val, (int, float)):
                    cell.number_format = "0"

        sl_cell = ws.cell(excel_row, 5)
        sl_cell.number_format = "0"

        formula_cell = ws.cell(excel_row, invoice.con_lai_col() + 1)
        formula_cell.value = invoice.con_lai_formula(r_idx)
        formula_cell.font = CELL_FONT
        formula_cell.fill = LOCKED_FILL
        formula_cell.alignment = CENTER
        formula_cell.number_format = "0"
        formula_cell.border = THIN

        ws.row_dimensions[excel_row].height = 18

    # --- Total Row ---
    total_row = HEADER_ROWS + len(invoice.rows) + 1
    ws.merge_cells(f"A{total_row}:D{total_row}")
    tot_label_cell = ws.cell(total_row, 1, "TỔNG CỘNG")
    tot_label_cell.font = Font(name="Calibri", bold=True, size=11)
    tot_label_cell.alignment = CENTER
    tot_fill = PatternFill("solid", fgColor="E2EFDA")
    
    for c in range(1, 5):
        cell_item = ws.cell(total_row, c)
        cell_item.border = THIN
        cell_item.fill = tot_fill
        cell_item.font = Font(name="Calibri", bold=True, size=11)

    first_data_r = HEADER_ROWS + 1
    last_data_r = HEADER_ROWS + len(invoice.rows)

    # Column 5: Total Quantity
    sl_letter = get_column_letter(invoice.sl_col() + 1)
    tot_sl_cell = ws.cell(total_row, 5, f"=SUM({sl_letter}{first_data_r}:{sl_letter}{last_data_r})")
    tot_sl_cell.font = Font(name="Calibri", bold=True, size=11)
    tot_sl_cell.alignment = CENTER
    tot_sl_cell.fill = tot_fill
    tot_sl_cell.border = THIN
    tot_sl_cell.number_format = "0"

    # Xe columns totals
    for xi in range(invoice.n_xe):
        col_idx = invoice.first_xe_col() + xi + 1
        c_letter = get_column_letter(col_idx)
        tot_xe_cell = ws.cell(total_row, col_idx, f"=SUM({c_letter}{first_data_r}:{c_letter}{last_data_r})")
        tot_xe_cell.font = Font(name="Calibri", bold=True, size=11)
        tot_xe_cell.alignment = CENTER
        tot_xe_cell.fill = tot_fill
        tot_xe_cell.border = THIN
        tot_xe_cell.number_format = "0"

    # Remaining column total
    cl_col_idx = invoice.con_lai_col() + 1
    cl_letter = get_column_letter(cl_col_idx)
    tot_cl_cell = ws.cell(total_row, cl_col_idx, f"=SUM({cl_letter}{first_data_r}:{cl_letter}{last_data_r})")
    tot_cl_cell.font = Font(name="Calibri", bold=True, size=11)
    tot_cl_cell.alignment = CENTER
    tot_cl_cell.fill = tot_fill
    tot_cl_cell.border = THIN
    tot_cl_cell.number_format = "0"

    ws.row_dimensions[total_row].height = 20

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = None
    ws.print_area = f"A1:{last_letter}{total_row}"

    ws.oddHeader.left.text = COMPANY_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


XE_RE = re.compile(r"^\s*xe\s*\(\s*(\d+)\s*\)\s*$", re.I)


def load_invoice(path: str | Path) -> Invoice:
    path = Path(path)
    wb = load_workbook(path, data_only=False)
    ws = wb.active

    header_row = None
    headers = []
    for r in range(1, 40):
        row_vals = [ws.cell(r, c).value for c in range(1, 80)]
        texts = [str(v).strip() if v is not None else "" for v in row_vals]
        if "STT" in texts and any(t.lower().startswith("mã") or t.lower().startswith("ma") for t in texts):
            header_row = r
            headers = texts
            break
    if header_row is None:
        raise ValueError("Không tìm thấy dòng tiêu đề STT / Mã hàng trong file.")

    def find_col(*names: str) -> int | None:
        lowered = [n.lower() for n in names]
        for i, h in enumerate(headers):
            if h.lower() in lowered:
                return i
        return None

    col_stt = find_col("stt") or 0
    col_ma = find_col("mã hàng", "ma hang", "mã") or 1
    col_ten = find_col("tên hàng", "ten hang", "tên") or 2
    col_dvt = find_col("đơn vị", "don vi", "dvt") or 3
    col_sl = find_col("số lượng", "so luong") or 4
    col_cl = find_col("còn lại", "con lai")

    xe_cols: list[tuple[int, int]] = []
    for i, h in enumerate(headers):
        m = XE_RE.match(h)
        if m:
            xe_cols.append((int(m.group(1)), i))
    xe_cols.sort()
    n_xe = max((n for n, _ in xe_cols), default=0)
    if n_xe == 0:
        # unnamed numeric columns between SL and Còn lại
        start = col_sl + 1
        end = col_cl if col_cl is not None else start
        n_xe = max(1, end - start)
        xe_index = list(range(start, start + n_xe))
    else:
        mapping = {n: i for n, i in xe_cols}
        xe_index = [mapping.get(i + 1, mapping[max(mapping)]) for i in range(n_xe)]

    rows: list[ProductRow] = []
    r = header_row + 1
    while r <= ws.max_row:
        cell1_str = str(ws.cell(r, 1).value or "").strip().upper()
        cell2_str = str(ws.cell(r, 2).value or "").strip().upper()
        cell3_str = str(ws.cell(r, 3).value or "").strip().upper()
        if any(term in cell1_str or term in cell2_str or term in cell3_str for term in ("TỔNG CỘNG", "TONG CONG", "TỔNG")):
            break
        ma = ws.cell(r, col_ma + 1).value
        ten = ws.cell(r, col_ten + 1).value
        sl = ws.cell(r, col_sl + 1).value

        if ma is None and ten is None and sl is None:
            # stop after a few empty rows
            if not rows:
                r += 1
                continue
            rest_empty = True
            for look in range(r, min(r + 3, ws.max_row + 1)):
                if any(ws.cell(look, c).value not in (None, "") for c in range(1, 8)):
                    rest_empty = False
                    break
            if rest_empty:
                break
            r += 1
            continue
        xe_vals: list[float | None] = []
        for idx in xe_index:
            xe_vals.append(parse_number(ws.cell(r, idx + 1).value))
        rows.append(
            ProductRow(
                ma="" if ma is None else str(ma).strip(),
                ten="" if ten is None else str(ten).strip(),
                dvt=str(ws.cell(r, col_dvt + 1).value or "Cái").strip(),
                so_luong=parse_number(sl) or 0.0,
                xe=xe_vals,
            )
        )
        r += 1

    customer = _read_labeled(ws, "tên khách") or ""
    address = _read_labeled(ws, "địa chỉ") or ""
    number = _read_labeled(ws, "số") or ""
    inv_date = _read_date(ws)

    inv = Invoice(
        customer_name=str(customer),
        customer_address=str(address),
        invoice_date=inv_date,
        invoice_number=str(number) if number is not None else "",
        n_xe=max(n_xe, 1),
        rows=rows or [ProductRow()],
        file_path=str(path),
    )
    return inv


def _read_labeled(ws, label: str):
    target = label.lower()
    for r in range(1, 16):
        for c in range(1, 40):
            val = ws.cell(r, c).value
            if val is not None:
                val_str = str(val).strip().rstrip(":").lower()
                if val_str == target:
                    for dc in range(1, 10):
                        nxt = ws.cell(r, c + dc).value
                        if nxt not in (None, ""):
                            return nxt
                    below = ws.cell(r + 1, c + 1).value
                    if below not in (None, ""):
                        return below
    return None


def _read_date(ws) -> date:
    for r in range(1, 16):
        for c in range(1, 40):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip().lower() == "ngày":
                for dc in range(1, 6):
                    nxt = ws.cell(r, c + dc).value
                    if isinstance(nxt, datetime):
                        return nxt.date()
                    if isinstance(nxt, date):
                        return nxt
                    parsed = _try_parse_date(nxt)
                    if parsed:
                        return parsed
    return date.today()


def _try_parse_date(value) -> date | None:
    if value is None:
        return None
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None
